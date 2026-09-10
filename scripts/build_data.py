import json
import os
from datetime import date, datetime
from openpyxl import load_workbook

SOURCE = os.environ.get('SOURCE_XLSX', 'Monitoramento Anual - 2026.xlsx')
SHEET = os.environ.get('SOURCE_SHEET', 'Base - BI')
OUTPUT = os.environ.get('OUTPUT_JSON', 'data.json')

wb = load_workbook(SOURCE, read_only=True, data_only=True)
ws = wb[SHEET]
rows = ws.iter_rows(values_only=True)
headers = next(rows)
headers = [str(h).strip() if h is not None else '' for h in headers]

def clean(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v

data = []
for row in rows:
    if not row or not any(v is not None and str(v).strip() for v in row):
        continue
    data.append({h: clean(v) for h, v in zip(headers, row) if h})

with open(OUTPUT, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, separators=(',', ':'))

print(f'{len(data)} registros exportados de {SHEET} para {OUTPUT}')
